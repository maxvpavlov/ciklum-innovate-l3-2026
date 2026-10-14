"""Build the 12 instructor guides (.docx and .md) and LEARNER_PROMPTS.md from specs/*.json.

Usage:
    uv run --with python-docx python build.py [--out-guides DIR] [--out-prompts FILE]
                                              [--template DOCX] [--lessons 1,5,12]

Defaults write to the real locations (instructor-guides/ and
incident-platform/LEARNER_PROMPTS.md). Derived facts are computed from git, not stored:
the tag hash, the code excerpts and the "Files changed" list.

    uv run --with python-docx python build.py --make-template SOURCE.docx
rebuilds template.docx (styles, numbering, page setup and footer of SOURCE, empty body).
"""

import argparse
from pathlib import Path

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from guidelib import (GUIDES_DIR, HERE, LESSONS, PROMPTS_FILE, SPECS_DIR, excerpt,
                      expected_json_text, files_changed, guide_filename, guide_title,
                      lesson_vars, load_common, load_spec, tag_hash, template)

from markdown import build_index, build_markdown

TEMPLATE = HERE / "template.docx"
FIGURES_DIR = HERE.parent / "console-screenshots"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
STYLE_IDS = {"Title": "Title", "Heading 1": "Heading1", "Heading 2": "Heading2",
             "Agent Prompt": "AgentPrompt", "Code": "Code", "List Bullet": "ListBullet"}
LINK_COLOR = "1155CC"
AGENDA_WIDTHS = [1440, 6696, 1944]      # twentieths of a point; 10080 = full text width
AGENDA_HEADER_FILL = "DCE6F1"
AGENDA_BORDER_COLOR = "D9D9D9"
# The original generator left the Activity and Format cells of "Agent build" rows without
# the compact paragraph spacing. Kept so the output stays byte-identical to the guides.
UNSPACED_AGENDA_FORMATS = {"Agent build"}


# ---------------------------------------------------------------- low-level XML

def el(tag, **attrs):
    """OxmlElement with attributes in the given order ("w_val" -> "w:val")."""
    return OxmlElement(tag, attrs={qn(k.replace("_", ":", 1)): str(v) for k, v in attrs.items()})


def text_el(text):
    t = el("w:t")
    t.text = text
    if text != text.strip():
        t.set(XML_SPACE, "preserve")
    return t


def run(text, bold=False, size=None, link=False):
    """A run. Newlines become <w:br/>; empty lines add no <w:t>, as python-docx does."""
    r = el("w:r")
    if bold or size or link:
        rpr = el("w:rPr")
        r.append(rpr)
        if bold:
            rpr.append(el("w:b"))
        if link:
            rpr.append(el("w:color", w_val=LINK_COLOR))
            rpr.append(el("w:u", w_val="single"))
        if size:
            rpr.append(el("w:sz", w_val=size))
    for i, line in enumerate(text.split("\n")):
        if i:
            r.append(el("w:br"))
        if line:
            r.append(text_el(line))
    return r


class GuideWriter:
    """Appends guide content to a document opened from the template."""

    def __init__(self, template_path):
        self.doc = docx.Document(str(template_path))
        self.body = self.doc.element.body
        self.sect_pr = self.body.find(qn("w:sectPr"))

    def add(self, element):
        self.sect_pr.addprevious(element)
        return element

    def para(self, style=None, runs=(), keep_lines=False, keep_next=False, page_break=False):
        p = el("w:p")
        if style or keep_lines or keep_next or page_break:
            ppr = el("w:pPr")
            if style:
                ppr.append(el("w:pStyle", w_val=STYLE_IDS[style]))
            if keep_next:
                ppr.append(el("w:keepNext"))
            if keep_lines:
                ppr.append(el("w:keepLines"))
            if page_break:
                ppr.append(el("w:pageBreakBefore"))
            p.append(ppr)
        for r in runs:
            p.append(r)
        return self.add(p)

    def text(self, text, style=None, **fmt):
        return self.para(style, [run(text)] if text else [], **fmt)

    def labelled(self, label, text, style=None, **fmt):
        return self.para(style, [run(label, bold=True), run(text)], **fmt)

    def link(self, text, url):
        rid = self.doc.part.relate_to(url, RT.HYPERLINK, is_external=True)
        h = el("w:hyperlink", r_id=rid)
        h.append(run(text, link=True))
        return h

    def code(self, text):
        return self.text(text, "Code", keep_lines=True)

    def bullets(self, items):
        for item in items:
            self.text(item, "List Bullet")

    def table(self, header, rows, widths):
        tbl = el("w:tbl")
        tbl_pr = el("w:tblPr")
        tbl_pr.append(el("w:tblW", w_type="auto", w_w=0))
        tbl_pr.append(el("w:jc", w_val="center"))
        tbl_pr.append(el("w:tblLayout", w_type="fixed"))
        tbl_pr.append(el("w:tblLook", w_firstColumn=1, w_firstRow=1, w_lastColumn=0, w_lastRow=0,
                         w_noHBand=0, w_noVBand=1, w_val="04A0"))
        borders = el("w:tblBorders")
        for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            borders.append(el(f"w:{side}", w_val="single", w_sz=4, w_color=AGENDA_BORDER_COLOR))
        tbl_pr.append(borders)
        tbl.append(tbl_pr)
        grid = el("w:tblGrid")
        for w in widths:
            grid.append(el("w:gridCol", w_w=w))
        tbl.append(grid)
        for is_header, cells in [(True, header)] + [(False, row) for row in rows]:
            tr = el("w:tr")
            if is_header:
                tr_pr = el("w:trPr")
                tr_pr.append(el("w:tblHeader"))
                tr.append(tr_pr)
            for col, (text, width) in enumerate(zip(cells, widths)):
                tc = el("w:tc")
                tc_pr = el("w:tcPr")
                tc_pr.append(el("w:tcW", w_type="dxa", w_w=width))
                if is_header:
                    tc_pr.append(el("w:shd", w_fill=AGENDA_HEADER_FILL))
                tc_pr.append(el("w:vAlign", w_val="center"))
                mar = el("w:tcMar")
                for side in ["top", "bottom", "left", "right"]:
                    mar.append(el(f"w:{side}", w_w=70, w_type="dxa"))
                tc_pr.append(mar)
                tc.append(tc_pr)
                p = el("w:p")
                if col == 0 or cells[-1] not in UNSPACED_AGENDA_FORMATS:
                    ppr = el("w:pPr")
                    ppr.append(el("w:spacing", w_after=40, w_line=240, w_lineRule="auto"))
                    p.append(ppr)
                p.append(run(text, bold=is_header, size=20))
                tc.append(p)
                tr.append(tc)
            tbl.append(tr)
        return self.add(tbl)

    def figure(self, image, caption, label):
        """A full-width console screenshot followed by a labelled caption."""
        picture = self.doc.add_paragraph()
        picture.add_run().add_picture(str(FIGURES_DIR / image), width=Inches(6.5))
        picture.paragraph_format.keep_with_next = True
        self.sect_pr.addprevious(picture._p)
        self.labelled(label, caption)

    def set_footer(self, text):
        footer = self.doc.sections[0].footer
        t = footer._element.find(".//" + qn("w:t"))
        t.text = text

    def save(self, path):
        self.doc.save(str(path))


# ---------------------------------------------------------------- guide content

def renderer(spec, common):
    """T(key, **extra): the lesson's version of a template, with all placeholders filled."""
    n = spec["number"]
    values = {**lesson_vars(n), "hash": tag_hash(n), "duration": spec["duration"],
              "title": guide_title(spec)}
    return lambda key, **extra: template(common, spec, key).format(**values, **extra)


def build_guide(spec, common, template_path, out_path):
    n = spec["number"]
    T = renderer(spec, common)
    L = common["labels"]
    H = common["headings"]
    g = GuideWriter(template_path)

    g.doc.core_properties.title = T("core_title")
    g.set_footer(T("footer"))

    # Header
    g.text(T("title"), "Title")
    g.text(T("subtitle"))
    g.labelled(L["outcome"], spec["outcome"])
    g.text(T("header"))

    # Where the session sits in the platform being built
    if "platform_role" in spec:
        g.text(H["platform_map"], "Heading 1")
        g.text(spec["platform_role"])
        status = lambda k: "Built earlier" if k < n else "This session" if k == n else "Later session"
        g.table(common["platform_map_header"],
                [[f"{k:02d}", name, status(k)] for k, name in common["platform_components"]], AGENDA_WIDTHS)
        g.text("")

    # Preparation
    g.text(H["preparation"], "Heading 1")
    for text in spec["preparation"]:
        g.text(text)
    g.para(runs=[run(T("access_before")), g.link(T("access_link"), T("access_url")),
                 run(T("access_after"))])
    g.text(T("agent_role"))
    g.labelled(L["prompt"], T("prepare_prompt"), "Agent Prompt", keep_lines=True)
    g.text(T("orient_intro"))
    g.labelled(L["prompt"], T("orient_prompt"), "Agent Prompt", keep_lines=True)
    g.labelled(L["readiness"], T("readiness"))
    for text in spec.get("readiness_notes", []):
        g.text(text)
    g.labelled(L["offline"], T("offline"))
    g.text(T("pinned"))

    # Agenda
    g.text(H["agenda"], "Heading 1", page_break=True)
    rows = spec.get("agenda") or common["agenda"][str(spec["duration"])]
    g.table(common["agenda_header"], rows, AGENDA_WIDTHS)
    g.text("")

    # Instructor script
    g.text(H["script"], "Heading 1")
    g.text(T("script_intro"))
    g.text(H["opening"], "Heading 2")
    g.labelled(L["say"], spec["opening_say"])
    g.labelled(L["show"], T("opening_show"))
    g.text(H["concept"], "Heading 2")
    g.labelled(L["say"], spec["concept_say"])
    g.labelled(L["pause"], T("concept_pause"))
    figures = spec.get("figures", [])
    for fig in [f for f in figures if f["after"] == "concept"]:
        g.figure(fig["image"], fig["caption"], L["figure"])

    for k, step in enumerate(spec["steps"], 1):
        ex = step["excerpt"]
        g.text(T("step_heading", k=k, step_title=step["title"]), "Heading 2")
        g.labelled(L["say"], step["say"])
        g.labelled(L["show"], T("step_show", file=step["file"]))
        g.labelled(L["prompt"], T("step_prompt", body=step["prompt"]), "Agent Prompt",
                   keep_lines=True)
        g.labelled(L["review"], T("review", start=ex["start"], end=ex["end"]), keep_next=True)
        g.code(excerpt(n, ex.get("file", step["file"]), ex["start"], ex["end"]))
        g.labelled(L["explain"], step["explain"])
        g.labelled(L["ask"], step["ask"])
        g.labelled(L["answer"], step["answer"])
        g.labelled(L["pause"], step.get("pause") or T("step_pause"))
        for fig in [f for f in figures if f["after"] == f"step{k}"]:
            g.figure(fig["image"], fig["caption"], L["figure"])

    # Learner practice
    g.text(H["practice_checkpoints"], "Heading 1")
    g.labelled(L["prompt"], T("exercise_prompt", body=spec["exercise"]), "Agent Prompt")
    g.labelled(L["verify"], T("verify"))
    g.labelled(L["prompt"], T("verify_prompt"), "Agent Prompt", keep_lines=True)
    g.labelled(L["expected"], T("expected_behaviour"))
    g.code(expected_json_text(spec))
    g.labelled(L["pause"], T("practice_pause"))
    for fig in [f for f in figures if f["after"] == "practice"]:
        g.figure(fig["image"], fig["caption"], L["figure"])

    # Completion criteria and files changed
    g.text(H["completion"], "Heading 1")
    g.bullets(spec["completion_criteria"])
    result = spec["recorded_test_result"]
    g.text(T("recorded", passed=result["passed"], seconds=result["seconds"]))
    g.text(H["files_changed"], "Heading 2")
    g.text(T("files_intro"))
    g.bullets(files_changed(n))
    g.labelled(L["prompt"], T("review_prompt"), "Agent Prompt", keep_lines=True)

    # Troubleshooting
    g.text(H["troubleshooting"], "Heading 1")
    for entry in spec["troubleshooting"]:
        g.text(entry["heading"], "Heading 2")
        g.labelled(L["prompt"], troubleshooting_prompt(entry, T), "Agent Prompt")

    # Practice, verification scope, close, references
    g.text(H["practice"], "Heading 1")
    g.bullets(spec["practice"])
    g.labelled(L["discussion"], T("discussion"))
    g.text(H["verification_scope"], "Heading 1")
    for item in spec["verification_scope"]:
        g.text(common["verification_scope"][item[1:]] if item.startswith("@") else item)
    g.text(H["close"], "Heading 1")
    g.labelled(L["say"], spec["close_say"])
    g.text(H["references"], "Heading 1")
    for ref in spec["references"]:
        g.para("List Bullet", [g.link(ref["title"], ref["url"])])

    g.save(out_path)


def troubleshooting_prompt(entry, T):
    if "prompt" in entry:
        return entry["prompt"]
    return T("recovery_prompt", symptom=entry.get("symptom", entry["heading"] + "."),
             diagnostic=entry["diagnostic"], recovery=entry["recovery"])


def lesson_prompts(spec, common):
    """The 13 prompts of a lesson, in order, as (markdown heading, prompt text)."""
    T = renderer(spec, common)
    S = common["learner_prompts"]["sections"]
    prompts = [(S["prepare"], T("prepare_prompt")), (S["orient"], T("orient_prompt"))]
    for k, step in enumerate(spec["steps"], 1):
        prompts.append((S["step"].format(k=k), T("step_prompt", body=step["prompt"])))
    prompts += [(S["exercise"], T("exercise_prompt", body=spec["exercise"])),
                (S["verify"], T("verify_prompt")),
                (S["review"], T("review_prompt"))]
    for entry in spec["troubleshooting"]:
        prompts.append((S["recover"].format(heading_lower=entry["heading"].lower()),
                        troubleshooting_prompt(entry, T)))
    return prompts


def build_prompts_md(specs, common):
    lp = common["learner_prompts"]
    parts = ["# " + lp["title"], *lp["intro"]]
    for spec in specs:
        v = lesson_vars(spec["number"])
        parts.append("## " + lp["lesson_heading"].format(**v, title=spec["title"]))
        parts.append(lp["references"].format(**v))
        for heading, prompt in lesson_prompts(spec, common):
            parts += ["### " + heading, "> " + prompt]
    return "\n\n".join(parts) + "\n"


# ---------------------------------------------------------------- template

def make_template(source, dest):
    """Copy a guide, keeping styles/sectPr/footer but removing body content and hyperlinks."""
    doc = docx.Document(str(source))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    rels = doc.part.rels
    for rid in [rid for rid, rel in rels.items() if rel.reltype == RT.HYPERLINK]:
        del rels[rid]
    doc.sections[0].footer._element.find(".//" + qn("w:t")).text = "Lesson 00    "
    doc.core_properties.title = "Lesson 00 instructor guide"
    doc.save(str(dest))


# ---------------------------------------------------------------- main

def build_all(specs_dir=SPECS_DIR, out_guides=GUIDES_DIR, out_prompts=PROMPTS_FILE,
              template_path=TEMPLATE, lessons=None, log=print):
    """Write the guides (and, for a full build, LEARNER_PROMPTS.md); return the written paths."""
    common = load_common()
    full = lessons is None
    specs = [load_spec(n, specs_dir) for n in (LESSONS if full else lessons)]
    out_guides = Path(out_guides)
    out_guides.mkdir(parents=True, exist_ok=True)
    written = []
    for spec in specs:
        path = out_guides / guide_filename(spec)
        build_guide(spec, common, template_path, path)
        md_path = path.with_suffix(".md")
        md_path.write_text(build_markdown(spec, common, renderer(spec, common), troubleshooting_prompt),
                           encoding="utf-8")
        written += [path, md_path]
        log(f"wrote {path} and {md_path.name}")
    if full:
        index = out_guides / "README.md"
        index.write_text(build_index(specs, common), encoding="utf-8")
        written.append(index)
        out_prompts = Path(out_prompts)
        out_prompts.parent.mkdir(parents=True, exist_ok=True)
        out_prompts.write_text(build_prompts_md(specs, common), encoding="utf-8")
        written.append(out_prompts)
        log(f"wrote {out_prompts}")
    return written


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--specs", default=str(SPECS_DIR), help="directory holding NN.json specs")
    ap.add_argument("--out-guides", default=str(GUIDES_DIR), help="directory for the .docx guides")
    ap.add_argument("--out-prompts", default=str(PROMPTS_FILE), help="path of LEARNER_PROMPTS.md")
    ap.add_argument("--template", default=str(TEMPLATE), help="style template .docx")
    ap.add_argument("--lessons", help="comma-separated lesson numbers (default: all 12; "
                                      "LEARNER_PROMPTS.md is written only for a full build)")
    ap.add_argument("--make-template", metavar="SOURCE", help="rebuild the template from a guide")
    args = ap.parse_args()

    if args.make_template:
        make_template(args.make_template, args.template)
        print(f"wrote {args.template}")
        return
    lessons = [int(x) for x in args.lessons.split(",")] if args.lessons else None
    build_all(args.specs, args.out_guides, args.out_prompts, args.template, lessons)


if __name__ == "__main__":
    main()
