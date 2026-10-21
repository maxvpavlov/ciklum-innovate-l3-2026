"""Build the Session 01 deck on the Ciklum template: new slides from its layouts, sample slides removed."""
import copy, sys
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
TEMPLATE_SLIDES = list(prs.slides._sldIdLst)
M = prs.slide_masters[1]
LAY = {l.name.strip(): l for l in M.slide_layouts}

NAVY, BLUE, MINT, PAPER = "011126", "0020BB", "00FFBE", "F9FAFB"
BRIGHT, CYAN, INK, GREY, LINE, SOFT = "143BFC", "50EAFF", "3C3C3C", "6B7280", "D5DAE3", "EEF1FF"
BODY, MED, MONO = "Zalando Sans Light", "Zalando Sans Medium", "IBM Plex Mono"
REPO = "github.com/maxvpavlov/ciklum-innovate-l3-2026"
PAGE = "maxvpavlov.github.io/ciklum-innovate-l3-2026/lesson-01"


def rgb(h): return RGBColor.from_string(h)


def slide(layout, title=None, notes=None):
    s = prs.slides.add_slide(LAY[layout])
    if title is not None:
        for ph in s.placeholders:
            if ph.placeholder_format.idx == 0:
                ph.text_frame.text = title
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def ph(s, idx, text):
    for p in s.placeholders:
        if p.placeholder_format.idx == idx:
            p.text_frame.text = text
            return p
    raise KeyError(idx)


def drop_empty_placeholders(s):
    for p in list(s.placeholders):
        if not p.has_text_frame or not p.text_frame.text.strip():
            if p.placeholder_format.idx != 0 or not p.text_frame.text.strip():
                p._element.getparent().remove(p._element)


def text(s, x, y, w, h, runs, size=12, color=INK, font=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space=4, margin=0.0):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, {opts})."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(o.get("size", size))
            f.name = o.get("font", font)
            f.color.rgb = rgb(o.get("color", color))
            if o.get("bold"): f.bold = True
            if o.get("italic"): f.italic = True
    return tb


def box(s, x, y, w, h, fill="FFFFFF", line=LINE, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=0.75):
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        b.adjustments[0] = min(0.5, radius / min(w, h))
    if fill: b.fill.solid(); b.fill.fore_color.rgb = rgb(fill)
    else: b.fill.background()
    if line: b.line.color.rgb = rgb(line); b.line.width = Pt(lw)
    else: b.line.fill.background()
    b.shadow.inherit = False
    b.text_frame.text = ""
    return b


def card(s, x, y, w, h, head, body, tag=None, fill="FFFFFF", head_color=BLUE, body_size=11):
    box(s, x, y, w, h, fill=fill)
    paras = [[(head, {"font": MED, "size": 13, "color": head_color})], body]
    if tag: paras.append([(tag, {"font": MONO, "size": 9, "color": BRIGHT})])
    text(s, x + 0.15, y + 0.12, w - 0.3, h - 0.2, paras, size=body_size, space=5)


def arrow(s, x1, y1, x2, y2, color=NAVY, width=1.25, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "dash")
    te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c


def code(s, x, y, w, h, src, size=9.5, hl=()):
    box(s, x, y, w, h, fill=NAVY, line=None, radius=0.1)
    lines = src.split("\n")
    paras = []
    for i, l in enumerate(lines):
        col = MINT if i in hl else "E6ECF5"
        paras.append([(l if l else " ", {"color": col})])
    t = text(s, x + 0.18, y + 0.14, w - 0.3, h - 0.2, paras, size=size, font=MONO, space=0)
    for p in t.text_frame.paragraphs:
        p.line_spacing = 1.0
    return t


def table(s, x, y, w, rows, widths, size=10.5, header=True, row_h=0.3, first_bold=True, hl_row=None):
    nr, nc = len(rows), len(rows[0])
    gt = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr)).table
    tbl = gt._tbl
    tp = tbl.tblPr
    for a in ("bandRow", "firstRow"): tp.set(a, "0")
    st = tp.find(qn("a:tableStyleId"))
    if st is not None: tp.remove(st)
    for i, cw in enumerate(widths): gt.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        gt.rows[r].height = Inches(row_h)
        for c, val in enumerate(row):
            cell = gt.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            is_head = header and r == 0
            fill = BLUE if is_head else ("E6FFF7" if hl_row == r else ("FFFFFF" if r % 2 else "F3F5F9"))
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(fill)
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.text = ""
            r0 = p.add_run(); r0.text = str(val)
            f = r0.font; f.size = Pt(size)
            f.name = MED if (is_head or (first_bold and c == 0)) else BODY
            mono = (not is_head) and str(val).startswith("`")
            if mono: r0.text = str(val).strip("`"); f.name = MONO; f.size = Pt(size - 1)
            f.color.rgb = rgb("FFFFFF" if is_head else NAVY)
    return gt


def chip(s, x, y, w, label, fill=SOFT, color=BLUE, size=10):
    b = box(s, x, y, w, 0.34, fill=fill, line=None, radius=0.17)
    t = text(s, x, y, w, 0.34, [[(label, {"font": MED, "size": size, "color": color})]], align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    return b


def subtitle(s, line, y=0.88):
    text(s, 0.4, y, 8.4, 0.5, line, size=13, color=INK)


PLAIN = "1_Plain slide with Confidential tag"
CONTENT = "2_Plain slide"

# ------------------------------------------------------------------ 1 cover
s = slide("Title Slide 1", "Building Agentic Platforms",
          "Welcome. This is Session 01 of twelve. Over the course you build one agentic platform, step by step, by "
          "directing your own coding agent. Today: why platforms, what they are made of, and the core every agent runs on: "
          "the ReAct runtime. In class I demonstrate; after class you repeat every prompt in your own project.")
t = ph(s, 0, "Building Agentic\vPlatforms")
for r in t.text_frame.paragraphs[0].runs: r.font.size = Pt(40)
ph(s, 1, "Innovate · L3 · Session 01 of 12 · 28 October 2026")
ph(s, 2, "Agent architecture and the ReAct runtime")

# ------------------------------------------------------------------ 2 course at a glance
s = slide(CONTENT, "The course at a glance",
          "Twelve sessions, twenty contact hours, one platform that you design, build, evaluate and operate. Everything "
          "runs offline on synthetic data with a scripted mock model, so nobody needs keys and everyone sees the same "
          "results. The capstone is your own tuned platform with a measured before/after comparison.")
stats = [("12", "sessions, one platform", "built step by step"), ("20", "contact hours", "plus practice after class"),
         ("1", "platform you build", "agents, tools, identity, console"), ("0", "API keys needed", "offline mock model")]
for i, (n, a, b) in enumerate(stats):
    x = 0.4 + i * 2.33
    box(s, x, 1.35, 2.13, 2.4)
    text(s, x + 0.2, 1.5, 1.8, 1.0, [[(n, {"font": "Zalando Sans ExtraLight", "size": 54, "color": BLUE})]])
    text(s, x + 0.2, 2.6, 1.8, 1.0, [[(a, {"font": MED, "size": 13, "color": NAVY})], b], size=11)
text(s, 0.4, 4.1, 9.0, 0.6, "From one agent loop to an orchestrated, governed platform: agents that use tools, delegate to "
     "other agents, talk to independent agent services, and run under identity, policy and evaluation you can show "
     "evidence for.", size=12)

# ------------------------------------------------------------------ 3 section
slide("Transition slide 2", "Why agentic platforms")

# ------------------------------------------------------------------ 4 agent teams
s = slide(CONTENT, "Beyond one agent",
          "An agent that reasons, calls a tool, reads the result and decides what to do next is no longer news. We still "
          "build that loop from scratch today, because everything later stands on it. If you know ReAct, treat today as "
          "the foundation you will keep extending. Many enterprise tasks are too big, too varied or too sensitive for one "
          "agent. Coding agents already show the pattern: an orchestrator plans a large feature, leads subagents with "
          "their own task and context, then reviews and combines what they return.")
subtitle(s, "Single agents are everyday tools. Enterprise work needs agents that work together.")
box(s, 3.35, 1.55, 3.3, 0.8, fill="E6EAFE", line=BLUE, lw=1.25)
text(s, 3.35, 1.6, 3.3, 0.75, [[("Orchestrator", {"font": MED, "size": 14, "color": NAVY})],
                                [("plans · delegates · reviews · combines", {"font": MONO, "size": 9, "color": GREY})]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
subs = [("Subagent: API change", "own context · fast model"), ("Subagent: data migration", "own context · DB access only"),
        ("Subagent: tests", "own context · runs in parallel")]
for i, (a, b) in enumerate(subs):
    x = 0.55 + i * 3.05
    box(s, x, 3.15, 2.85, 0.85)
    text(s, x, 3.2, 2.85, 0.8, [[(a, {"font": MED, "size": 12, "color": NAVY})],
                                [(b, {"font": MONO, "size": 9, "color": GREY})]], align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    arrow(s, 5.0, 2.35, x + 1.425, 3.12)
text(s, 0.4, 4.25, 9.0, 0.4, "The orchestration pattern from coding agents. The same shape appears in support, operations, "
     "finance and research work.", size=10.5, color=GREY, align=PP_ALIGN.CENTER)

# ------------------------------------------------------------------ 5 when to split
s = slide(CONTENT, "When do you need more agents?",
          "Four reasons, each tied to the part of the course that builds it. And the counterweight: if none applies, one "
          "agent is the right design, because every agent adds cost, latency and failure modes.")
items = [("The work outgrows one context", "It will not fit one context window, and step-by-step compaction does not help "
          "because the work splits into parallel parts that each need their own context.", "supervisor and subagents"),
         ("It is better done in a clean context", "Even when it fits, a sub-task goes better in a fresh context without "
          "unrelated history, other branches of the work or earlier mistakes.", "scoped subagents"),
         ("Each sub-task needs a different model", "A strong model plans and judges; faster, cheaper models handle routine "
          "sub-tasks. Switching changes quality, cost and latency.", "model routing"),
         ("Access must be isolated", "An agent is bound to one system or service, or acts for one person, and sees only "
          "that system's data and that person's permissions.", "A2A · gateway · identity")]
for i, (h, b, t) in enumerate(items):
    x, y = 0.4 + (i % 2) * 4.65, 1.05 + (i // 2) * 1.9
    card(s, x, y, 4.45, 1.72, h, b, t)

# ------------------------------------------------------------------ 6 quote
s = slide("Quote Slide 2", "If none of these applies, one agent is the right design. Every agent you add has to earn its place.",
          "Say it plainly: more agents are not better by default. Each one adds cost, latency and failure modes.")

# ------------------------------------------------------------------ 7 why a platform
s = slide(CONTENT, "A platform is the natural next step",
          "Once you run many agents, something has to spawn them, configure them, host them and give everyone one place to "
          "manage what they do. That is the agentic platform: the central hub for agentic operations. It adds the "
          "traceability and control points a single agent cannot provide on its own. Walk through the six questions.")
subtitle(s, "Something has to spawn, configure, host and manage the agents, and answer six questions for every run.")
qs = [("Who called what?", "Every tool call passes one policy check and lands in one audit record.", "tool gateway · audit"),
      ("On whose behalf?", "Each agent acts for a known person, with narrow, short-lived credentials.", "identity · approvals"),
      ("When does it stop?", "Every run has step, token and time limits.", "bounded runtime"),
      ("Is version 2 better?", "Changes ship on measured quality, cost and latency.", "evaluation · release gates"),
      ("How do agents work together?", "A supervisor delegates inside; A2A reaches other teams' agents.", "orchestration · A2A"),
      ("What actually happened?", "Which version ran, every decision and every failure.", "traces · registry · console")]
for i, (h, b, t) in enumerate(qs):
    x, y = 0.4 + (i % 3) * 3.1, 1.55 + (i // 3) * 1.68
    card(s, x, y, 2.95, 1.5, h, b, t, body_size=10.5)

# ------------------------------------------------------------------ 8 parts table
s = slide(CONTENT, "The main parts of a platform",
          "This is the map of the course. Today we build the first row, the agent runtime: the library every agent runs on. "
          "Everything else wraps it rather than replacing it.")
rows = [["Part", "The question it answers"],
        ["Agent runtime · today", "How does one agent act, and how is it forced to stop?"],
        ["Tools and MCP", "How do agents use real systems through reliable, permissioned contracts?"],
        ["Context and memory", "What does the agent know, from which sources, and for how long?"],
        ["Evaluation", "Is it actually good, and better than the last version?"],
        ["Durability and approval", "What survives a crash, and who signs off before an effect happens?"],
        ["Orchestration", "How does a supervisor split work across scoped subagents?"],
        ["Agent-to-agent (A2A)", "How do independent agent services discover and call each other?"],
        ["Registry and console", "Which agents exist, which version ran, and what happened?"],
        ["Tool gateway", "Is every tool call checked against one policy and recorded?"],
        ["Identity and credentials", "On whose behalf does each agent act, and with which secrets?"],
        ["Operations and releases", "Can we trace, throttle, roll out and roll back safely?"],
        ["Capstone", "Can you show a measured improvement and defend your design?"]]
table(s, 0.4, 0.95, 9.0, rows, [2.6, 6.4], size=9.5, row_h=0.31, hl_row=1)

# ------------------------------------------------------------------ 9 blueprint
s = slide(CONTENT, "The platform we build",
          "The highlighted box is today's work. The console across the top, the registry and supervisor on the left, the "
          "gateway and identity on the right, and context, evaluation, durability and operations underneath all wrap the "
          "runtime rather than replacing it.")
box(s, 0.6, 0.95, 8.8, 0.7)
text(s, 0.8, 1.0, 8.4, 0.62, [[("Platform console", {"font": MED, "size": 13, "color": NAVY})],
                              [("registry · runs · approvals · gateway audit · identity · operations", {"font": MONO, "size": 9, "color": GREY})]])
for x, h1, b1, h2, b2 in [(0.6, "Agent registry", "versioned agents, remote agents, tool servers", "Supervisor + A2A", "subagents, remote specialists"),
                          (6.9, "MCP tool gateway", "policy and audit on every tool call", "Identity", "delegated, short-lived credentials")]:
    box(s, x, 1.95, 2.5, 2.0)
    text(s, x + 0.18, 2.05, 2.15, 1.85, [[(h1, {"font": MED, "size": 13, "color": NAVY})], [(b1, {"size": 10, "color": GREY})],
                                         " ", [(h2, {"font": MED, "size": 13, "color": NAVY})], [(b2, {"size": 10, "color": GREY})]])
box(s, 3.55, 2.05, 2.9, 1.8, fill="E6EAFE", line=BLUE, lw=1.5)
text(s, 3.75, 2.15, 2.5, 1.65, [[("Agent runtime", {"font": MED, "size": 14, "color": NAVY})],
                                [("bounded ReAct loop", {"font": MONO, "size": 9, "color": GREY})],
                                [("run contract · limits · events", {"font": MONO, "size": 9, "color": GREY})], " ",
                                [("TODAY", {"font": MED, "size": 11, "color": BLUE})]])
box(s, 0.6, 4.2, 8.8, 0.7)
text(s, 0.8, 4.25, 8.4, 0.62, [[("Context · evaluation · durability · operations", {"font": MED, "size": 13, "color": NAVY})],
                               [("memory · eval harness · checkpoints and approval · traces and releases · tuning", {"font": MONO, "size": 9, "color": GREY})]])
arrow(s, 3.1, 2.95, 3.52, 2.95); arrow(s, 6.9, 2.95, 6.48, 2.95)
arrow(s, 5.0, 2.05, 5.0, 1.68); arrow(s, 5.0, 3.85, 5.0, 4.17)

# ------------------------------------------------------------------ 10 how we work
s = slide(CONTENT, "How we work in this course",
          "You direct a coding agent; it types the code and runs the checks; you judge the evidence. Every session ends at a "
          "tagged, runnable checkpoint. A coding agent saying done proves nothing: we look at diffs, test output and run "
          "events. Everything is offline and synthetic.")
ws = [("You direct a coding agent", "You set scope, read the change and judge the evidence. The agent types the code and runs the checks."),
      ("Every session ends runnable", "Each session finishes at a checkpoint with tests that pass, pushed to your GitLab repository."),
      ("Evidence, not claims", "A coding agent saying \"done\" proves nothing. We look at diffs, test output and run events."),
      ("Offline and synthetic", "A deterministic mock model and invented data: same behaviour for everyone, no keys, no cost.")]
for i, (h, b) in enumerate(ws):
    x, y = 0.4 + (i % 2) * 4.65, 1.05 + (i // 2) * 1.75
    card(s, x, y, 4.45, 1.55, h, b)
text(s, 0.4, 4.6, 9.0, 0.35, "Capstone: your own tuned platform, a runbook, a live demo, and a before/after comparison of "
     "quality, cost and latency.", size=11, color=GREY)

# ------------------------------------------------------------------ 11 section
slide("Transition slide 2", "The premise")

# ------------------------------------------------------------------ 12 how it starts
s = slide(CONTENT, "The platform investigates incidents",
          "One invented domain throughout: an incident-investigation platform for two fictional companies, Astro and "
          "Locomotive. Picture how it starts: at Astro the payments team notices checkout is slow and reports it to the "
          "service desk, which opens a ticket. An investigation agent picks it up. It can only read, using two tools, and "
          "must answer with a cause and the evidence for it. A runbook is a document for people, not a script: the agent "
          "reads it as evidence and never runs it. A suggested fix such as a rollback stays text until a person approves it.")
subtitle(s, "The payments team reports slow checkout. An investigation agent picks up the ticket, reads, and must cite its evidence.")
cs = [("Ticket", "The report from the business: which service, what people noticed, and an ID such as INC-001. It is the agent's task."),
      ("Log record", "What monitoring captured. log_lookup returns the service, the signal (a 900 ms database wait after a deployment) and a suspected cause."),
      ("Runbook", "A team's written procedure for a known failure: how to recognise it, its usual cause, what to do. Read as evidence, never executed.")]
for i, (h, b) in enumerate(cs):
    card(s, 0.4 + i * 3.1, 1.6, 2.95, 2.3, h, b)
text(s, 0.4, 4.15, 9.0, 0.5, "A good teaching case: tools, evidence, tenants that must not see each other's data, approvals "
     "before any effect, and more than one agent. All data are synthetic; no real system is touched.", size=11, color=GREY)

# ------------------------------------------------------------------ 13 tickets
s = slide(CONTENT, "Astro's tickets, and case one",
          "Three tickets. INC-001 is today's case: payments got slow after a deployment. INC-002 is clock skew. INC-003 has "
          "no diagnostic logs, so the honest answer is not enough evidence. Keep INC-003 in mind for the question later.")
table(s, 0.4, 1.0, 9.0, [["Ticket", "Service", "What the logs show", "What a good agent concludes"],
                         ["INC-001 · today", "payments", "DB pool wait 900 ms after deploy d42", "connection pool exhaustion, with a matching runbook"],
                         ["INC-002", "login", "token expiry errors; host clock +180 s", "clock skew, with a matching runbook"],
                         ["INC-003", "payments", "latency elevated; no diagnostic logs", "\"not enough evidence\": nothing supports a cause"]],
      [1.5, 1.1, 3.0, 3.4], size=10, row_h=0.42, hl_row=1)
box(s, 0.4, 3.0, 9.0, 1.7)
text(s, 0.6, 3.1, 8.6, 1.55, [[("INC-001 · tenant astro · payments latency", {"font": MONO, "size": 10.5, "color": BLUE})],
                              [("Reported by  ", {"font": MED}), ("the payments team, via the service desk", {})],
                              [("Symptom  ", {"font": MED}), ("checkout is slow since this morning's deployment", {})],
                              [("Agent can use  ", {"font": MED}), ("log_lookup and runbook_search, read-only, Astro's data only", {})],
                              [("We want  ", {"font": MED}), ("a likely cause, with the evidence it came from", {})]], size=11.5, space=3)

# ------------------------------------------------------------------ 14 acceptance
s = slide(CONTENT, "What makes a good investigation?",
          "Ask the room before showing the three answers. Then: a fluent paragraph is not an acceptance test. Our first "
          "checkpoint must also stop when it runs out of allowance or starts repeating itself.")
crit = [("A likely cause", "stated plainly, not hedged into nothing."),
        ("The source of that claim", "the log or runbook IDs it actually read."),
        ("An honest \"not enough evidence\"", "when the data does not support a cause.")]
for i, (h, b) in enumerate(crit):
    x = 0.4 + i * 3.1
    box(s, x, 1.3, 2.95, 1.9)
    text(s, x + 0.2, 1.4, 0.8, 0.6, [[(f"0{i+1}", {"font": "Zalando Sans ExtraLight", "size": 30, "color": BLUE})]])
    text(s, x + 0.2, 2.05, 2.6, 1.1, [[(h, {"font": MED, "size": 13, "color": NAVY})], b], size=11)
box(s, 0.4, 3.55, 9.0, 0.9, fill=NAVY, line=None)
text(s, 0.6, 3.55, 8.6, 0.9, [[("A fluent paragraph is not an acceptance test.", {"font": BODY, "size": 18, "color": "FFFFFF"})]],
     anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------------------ 15 section
slide("Transition slide 2", "The ReAct loop")

# ------------------------------------------------------------------ 16 loop diagram
s = slide(CONTENT, "Reason, act, observe",
          "Today we build only the centre box of the platform. ReAct describes an agent as a repeating cycle: a thought "
          "about what to do next, an action, and an observation of the result, until it can answer. With native tool "
          "calling: the model's turn contains its reasoning and a tool request; the runtime runs the tool; the result goes "
          "back as an observation. The model never calls code. It proposes; the runtime decides whether and how to execute. "
          "That boundary is where authorization, limits and auditing live for the rest of the course.")
for x, h, b, hl in [(0.5, "Model", "reasons and proposes\na tool call or an answer", False),
                    (3.75, "Runtime", "checks limits, permits\nrecords the pending call\nexecutes, appends result", True),
                    (7.0, "Tools", "log_lookup\nrunbook_search", False)]:
    box(s, x, 1.2, 2.5, 1.75, fill="E6EAFE" if hl else "FFFFFF", line=BLUE if hl else LINE, lw=1.5 if hl else 0.75)
    text(s, x + 0.18, 1.3, 2.2, 1.6, [[(h, {"font": MED, "size": 14, "color": NAVY})]] +
         [[(l, {"font": MONO, "size": 9.5, "color": GREY})] for l in b.split("\n")], space=3)
arrow(s, 3.02, 1.75, 3.72, 1.75); arrow(s, 6.27, 1.75, 6.97, 1.75)
arrow(s, 6.97, 2.45, 6.27, 2.45); arrow(s, 3.72, 2.45, 3.02, 2.45)
for x, y, l in [(3.0, 1.42, "proposal"), (6.25, 1.42, "execute"), (6.25, 2.52, "observation"), (3.0, 2.52, "next turn")]:
    text(s, x - 0.15, y, 1.05, 0.3, [[(l, {"font": MONO, "size": 7.5, "color": GREY})]], align=PP_ALIGN.CENTER)
box(s, 0.5, 3.35, 9.0, 1.15, fill=PAPER)
text(s, 0.7, 3.45, 8.6, 1.0, [[("Thought", {"font": MED, "color": BLUE}), (" the model's reasoning in its turn   ", {}),
                               ("Action", {"font": MED, "color": BLUE}), (" a tool call the runtime may run   ", {}),
                               ("Observation", {"font": MED, "color": BLUE}), (" the tool result, sent back", {})],
                              [("The model never calls code. It proposes; the runtime decides. That boundary is where authorization, limits and auditing live.", {"color": GREY})]],
     size=11.5, space=6)

# ------------------------------------------------------------------ 17 turn table
s = slide(CONTENT, "What the model can send back",
          "Every call returns one assistant turn, and the provider reports why the turn ended. The runtime decides what "
          "happens next from that structure, never from the wording. Reasoning, visible text or a provider's separate "
          "thinking, is never an action and never evidence. A reply with no tool call that is not the answer must be handled "
          "explicitly, never guessed at. Be honest: our first runtime handles only the first two rows; cut-off, refused and "
          "non-final replies fail without a named reason until we harden it for real models.")
table(s, 0.4, 1.0, 9.0, [["The turn contains", "Ended because (OpenAI / Anthropic)", "What the runtime does"],
                         ["One or more tool calls", "`tool_calls / tool_use", "Runs the permitted calls, appends each observation with the same call ID, calls the model again"],
                         ["Text only", "`stop / end_turn", "A candidate final answer: checks shape and that every citation was observed"],
                         ["Text that was cut off", "`length / max_tokens", "Not an answer: stop with a named reason, or continue with a larger budget"],
                         ["A refusal or filtered output", "`content_filter / refusal", "Not an answer: stop with a named reason"]],
      [2.2, 2.4, 4.4], size=10, row_h=0.55)
text(s, 0.4, 4.0, 9.0, 0.7, [[("Reasoning is never an action and never evidence. ", {"font": MED, "color": NAVY}),
                              ("Today's runtime handles the first two rows; the others fail without a named reason until we harden it for real models.", {"color": GREY})]],
     size=11)

# ------------------------------------------------------------------ 18 agent vs sequence
s = slide(CONTENT, "Agent vs hard-coded sequence",
          "An AI investigation platform is worth building only where incidents are too varied to follow a path written in "
          "advance: the next step depends on what the last one revealed, evidence is spread across systems, and the "
          "investigation should end in a fix a person approves. Our tickets are deliberately small. INC-001 alone could be "
          "a hard-coded sequence; INC-002 needs a different runbook question once the log is read; INC-003 should end with "
          "an honest not enough evidence. Every change is measured against a baseline on the same cases, starting with "
          "this single agent.")
for x, h, sub, steps, foot, dark in [(0.4, "A hard-coded sequence", "The developer fixes every step in advance.",
                                      ["read the log", "search the runbook", "write the summary"],
                                      "Predictable, but it breaks on the first case that does not follow the script.", False),
                                     (5.05, "An agent-driven investigation", "The next permitted step is chosen from what was just observed.",
                                      ["propose → run → observe", "propose → run → observe", "… then answer, or stop"],
                                      "Adapts to each case, so it needs limits, cited evidence and approval before any fix.", True)]:
    box(s, x, 1.05, 4.5, 3.6, fill=NAVY if dark else "FFFFFF", line=None if dark else LINE)
    c1, c2 = ("FFFFFF", "C9D3F5") if dark else (NAVY, GREY)
    paras = [[(h, {"font": MED, "size": 15, "color": c1})], [(sub, {"color": c2})], " "]
    paras += [[(f"{i+1}. {t}", {"font": MONO, "size": 11, "color": MINT if dark else BLUE})] for i, t in enumerate(steps)]
    paras += [" ", [(foot, {"color": c1})]]
    text(s, x + 0.25, 1.2, 4.0, 3.35, paras, size=11.5, space=4)

# ------------------------------------------------------------------ 19 real run
s = slide(CONTENT, "One run of INC-001, step by step",
          "The actual run of the finished Session 01 code with the mock model. Three parties take turns: the model proposes, "
          "the runtime checks, executes and carries results back, the tool returns data. Trace the call ID from the "
          "proposal to its observation. The model runs nothing; it asks. The final turn has no tool call, so it is the "
          "answer, and the runtime checks both cited IDs really appeared in tool results.")
run = [["Who", "What happens"],
       ["Runtime", "Sends the instructions, the task \"Investigate INC-001\" and the two tool definitions"],
       ["Model · 1", "`tool_call call-0  log_lookup {\"incident_id\": \"INC-001\"}"],
       ["Runtime → Tool", "Checks permission and budget, records call-0 as pending, runs it"],
       ["Tool", "`call-0 → log-astro-001 · payments · DB pool wait 900ms after deploy d42"],
       ["Model · 2", "`tool_call call-1  runbook_search {\"query\": \"connection_pool_exhaustion\"}"],
       ["Tool", "`call-1 → rb-pool-v1 · \"Consider rollback after approval\""],
       ["Model · 3", "`{\"status\": \"complete\", \"cause\": \"connection_pool_exhaustion\", \"evidence\": [\"log-astro-001\", \"rb-pool-v1\"]}"],
       ["Runtime", "Checks the cited IDs were observed and marks the run completed; the rollback stays a suggestion"]]
table(s, 0.4, 0.95, 9.0, run, [1.45, 7.55], size=10, row_h=0.39)
text(s, 0.4, 4.6, 9.0, 0.3, [[("3 model steps · 2 tool calls · 75 fixture tokens · cost 0 (no prices configured)", {"font": MONO, "size": 10, "color": BLUE})]])

# ------------------------------------------------------------------ 20 completed vs solved
s = slide("Quote Slide 2", "Does a completed run mean the incident is solved?",
          "Ask, then answer: no. INC-003 also completes: the log has no diagnostic detail and no runbook matches, so the "
          "agent answers with status insufficient_evidence and cause unknown. A run finishing and its conclusion being "
          "supported by evidence are different checks.")

# ------------------------------------------------------------------ 21 ownership
s = slide(CONTENT, "The application owns the limits",
          "The model owns nothing. Tenant, run, role, permitted tools and limits are set by application code; model output "
          "is untrusted input that the runtime validates before acting on it. Events record what happened step by step; "
          "later the console shows them as a run timeline.")
table(s, 0.4, 1.0, 9.0, [["Field", "Owner", "Why it matters"],
                         ["`tenant, run_id, role", "application", "A retrieved document must never be able to switch tenant or identity"],
                         ["`allowed tools", "application", "Permissions are granted by code, never claimed in model output"],
                         ["`Limits: steps, tokens, seconds", "application", "Each limit catches a different failure; together they bound every run"],
                         ["`messages, tool calls, answer", "model", "Untrusted input: the runtime validates it before acting on it"],
                         ["`events", "runtime", "What happened, step by step; later the console's run timeline"]],
      [2.7, 1.3, 5.0], size=10.5, row_h=0.5)
text(s, 0.4, 4.25, 9.0, 0.4, [[("The model owns nothing.", {"font": MED, "size": 14, "color": BLUE})]])

# ------------------------------------------------------------------ 22 guards
s = slide(CONTENT, "Four guards: every stop has a name",
          "Each guard catches a different failure. step_budget stops a loop that never converges. token_budget reserves room "
          "for the reply before each call and rechecks reported usage. TimeoutError catches a hang that consumes no steps "
          "and no tokens. repeated_action stops an agent stuck asking the same question. Every stop is a named, explained "
          "outcome, never an unexplained exception. Poll: a provider hangs and returns nothing; which component stops the "
          "run? The runtime, with its elapsed-time limit.")
gs = [("step_budget", "default 8 model steps", "Stops a loop that never converges."),
      ("token_budget", "default 8,000, checked before each call", "Reserves room for the reply, then rechecks reported usage."),
      ("TimeoutError", "default 20 s for the whole run", "Catches a hang that consumes no steps and no tokens."),
      ("repeated_action", "same tool, same arguments", "Stops an agent stuck asking the same question.")]
for i, (c, l, d) in enumerate(gs):
    x = 0.4 + i * 2.33
    box(s, x, 1.2, 2.13, 2.6)
    text(s, x + 0.15, 1.35, 1.85, 2.4, [[(c, {"font": MONO, "size": 13, "color": BLUE})],
                                        [(l, {"font": MONO, "size": 9, "color": GREY})], " ", d], size=11.5, space=5)
text(s, 0.4, 4.1, 9.0, 0.5, [[("Poll: a provider hangs and returns nothing. Which component stops the run? ", {"font": MED, "color": NAVY}),
                              ("The runtime, with its elapsed-time limit.", {"color": GREY})]], size=11.5)

# ------------------------------------------------------------------ 23 section
slide("Transition slide 2", "How you build it")

# ------------------------------------------------------------------ 24 route
s = slide(CONTENT, "How you work through the build",
          "In class, watch me run every prompt on a fresh project. After class, run the same prompts, in the same order, in "
          "your own project. You do not type the code: each step is one prompt to your coding agent, which makes the change "
          "and runs the checks; you judge the result. A coding agent saying done is not evidence: look at the diff and at "
          "the check output. All prompts are on the session page and in LEARNER_PROMPTS.md in the course repository.")
route = ["Set up", "Add the check pack", "Build 1", "Build 2", "Build 3", "Build 4", "Break it", "Verify and push", "Review"]
widths = [0.8, 1.5, 0.72, 0.72, 0.72, 0.72, 0.85, 1.25, 0.75]
x = 0.4
for i, (r, w) in enumerate(zip(route, widths)):
    chip(s, x, 1.25, w, r, fill=BLUE if i == 0 else SOFT, color="FFFFFF" if i == 0 else BLUE, size=9)
    x += w + 0.06
links = [("Course repository", REPO), ("Session 01 kit", "learner-kit/lessons/01"),
         ("Every session's prompts", "LEARNER_PROMPTS.md"), ("Session page", PAGE)]
for i, (h, u) in enumerate(links):
    x = 0.4 + i * 2.33
    box(s, x, 2.05, 2.13, 1.2)
    text(s, x + 0.15, 2.15, 1.85, 1.05, [[(h, {"font": MED, "size": 12, "color": NAVY})], [(u, {"font": MONO, "size": 8.5, "color": BRIGHT})]], space=5)
text(s, 0.4, 3.6, 9.0, 0.9, [[("Rule for the whole build: ", {"font": MED, "color": NAVY}),
                              ("a coding agent saying \"done\" is not evidence. Look at the diff and at the check output.", {})]], size=12)

# ------------------------------------------------------------------ 25 setup
s = slide(CONTENT, "Set up: the lesson-00 prompt",
          "Create an empty folder such as ~/agentic-course, start your coding agent in it, and paste the lesson-00 prompt "
          "from the session page. It clones the public course repository as a read-only copy, creates your project from the "
          "kit's starter, installs everything, runs the first test, installs glab, signs you in to GitLab and pushes your "
          "project to a new private repository in the course group. Start your agent in the same folder every later session.")
steps = [("1", "Create an empty folder", "~/agentic-course, and start your coding agent in it"),
         ("2", "Paste the lesson-00 prompt", "from the session page; it does the rest"),
         ("3", "Expect", "1 passed · LESSON = 0 · your project in your GitLab group")]
for i, (n, h, b) in enumerate(steps):
    y = 1.1 + i * 1.05
    box(s, 0.4, y, 0.55, 0.55, fill=BLUE, line=None, shape=MSO_SHAPE.OVAL)
    text(s, 0.4, y, 0.55, 0.55, [[(n, {"font": MED, "size": 14, "color": "FFFFFF"})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.1, y - 0.02, 3.4, 0.9, [[(h, {"font": MED, "size": 13, "color": NAVY})], b], size=11)
code(s, 4.8, 1.1, 4.6, 2.6, "~/agentic-course/\n  course/          read-only copy of\n                   the course repository\n"
     "    learner-kit/\n      starter/     copied once\n      lessons/01/  given files and\n                   check pack\n"
     "  my-platform/     your project, your agent\n                   writes here, pushed to GitLab", size=9.5, hl=(7, 8))

# ------------------------------------------------------------------ 26 what you build
s = slide(CONTENT, "What you build today",
          "core.py is the agent runtime: the core of every agent you will run. It holds the data the application owns, the "
          "shape of a model's reply and of a run, and the ReAct loop. From Session 08 the platform hosts it; a supervisor "
          "built from the same loop coordinates other agents from Session 06. The given files are data and a scripted mock "
          "model everyone shares. In production agents run in separate processes or sandboxes, apart from the platform's "
          "control plane; in this course they run in the platform's own process, a compromise that keeps everything "
          "runnable on a laptop.")
code(s, 0.4, 1.0, 5.2, 3.6, "my-platform/\n  incident/\n    __init__.py   LESSON = 0, then 1\n    lab.py        runs a session's demo\n"
     "    core.py       you build: the runtime\n    providers.py  you build; mock is given\n    fixtures.py   given: synthetic data\n"
     "    labs/l01.py   check pack: the demo\n  tests/\n    test_setup.py\n    test_lesson01.py  check pack", size=10, hl=(4, 5))
text(s, 5.9, 1.0, 3.6, 3.6, [[("core.py is the agent runtime", {"font": MED, "size": 14, "color": NAVY})],
                             "The data the application owns, the shape of a model's reply and of a run, and the ReAct loop. "
                             "Every later session builds on it.",
                             [("Given files", {"font": MED, "size": 13, "color": NAVY})],
                             "Synthetic data and a scripted mock model everyone shares, so results are comparable.",
                             [("A teaching compromise", {"font": MED, "size": 13, "color": NAVY})],
                             "Here the platform runs agents in its own process; in production they run in separate processes or sandboxes."],
     size=11, space=6)

# ------------------------------------------------------------------ 27 step 1
s = slide(CONTENT, "Build step 1: the run contract",
          "Step 1 creates core.py with the data contracts only: Context, Limits, Reply and Run. Separate what the "
          "application owns from what the model produces. None of these fields come from the model. Ask: why not let the "
          "model's answer say which tools it may use? Because model output is untrusted; the application builds Context.")
code(s, 0.4, 1.0, 4.9, 2.35, "@dataclass\nclass Context:\n    tenant: str = \"astro\"\n    run_id: str = field(default_factory=...)\n"
     "    role: str = \"investigator\"\n    allowed: tuple = (\"log_lookup\",\n                      \"runbook_search\")\n"
     "    approved: tuple = ()\n    parent_id: str | None = None", size=10)
fields = [("tenant", "whose data the run may see"), ("run_id", "ties events and child runs together"),
          ("role", "what kind of actor; later, who may write"), ("allowed", "the only tools this run may call"),
          ("approved", "approved actions; empty until approvals"), ("parent_id", "the delegating run; empty until subagents")]
text(s, 5.6, 1.0, 3.9, 2.6, [[(f, {"font": MONO, "size": 10, "color": BLUE}), ("  " + d, {})] for f, d in fields], size=11, space=6)
box(s, 0.4, 3.65, 9.0, 0.9, fill=PAPER)
text(s, 0.6, 3.7, 8.6, 0.8, [[("Also in step 1: ", {"font": MED, "color": NAVY}),
                              ("Limits(steps, tokens, seconds, max_output_tokens) · Reply(message, tokens, cost_usd) · "
                               "Run(messages, status, steps, tokens, pending, answer, error)", {"font": MONO, "size": 10})],
                             [("None of these come from the model.", {"font": MED, "color": BLUE})]], size=11, space=4)

# ------------------------------------------------------------------ 28 whole loop
s = slide(CONTENT, "Step 3: the whole loop",
          "Steps 2 to 4 build the providers, this loop and its guards. Read the loop by its numbers: 1, 2, 4 and 8 are the "
          "limits; 5 is the thought, one model turn; 6 and 7 classify that turn: tool calls become pending work and go back "
          "to 3; anything else must be the final answer; 3 is the action and observation. Abridged: the real function also "
          "records events, pauses for approval and handles cancellation.")
loop = ("① async with asyncio.timeout(limits.seconds):\n②     while run.steps < limits.steps:\n"
        "③         while run.pending: ...act: run the tools...\n④         if tokens + estimate + max_out > limits.tokens:\n"
        "              raise ValueError(\"token_budget\")\n⑤         reply = await provider.complete(messages, schemas, max_out)\n"
        "          run.steps += 1;  run.messages.append(reply.message)\n⑥         if reply.message.get(\"tool_calls\"):\n"
        "              run.pending = [...];  save();  continue\n⑦         answer = json.loads(reply.message[\"content\"])\n"
        "          ...check shape and citations...;  return run\n⑧     raise ValueError(\"step_budget\")")
code(s, 0.4, 1.0, 6.2, 3.7, loop, size=9.5, hl=(2,))
text(s, 6.85, 1.0, 2.65, 3.7, [[("① ② ④ ⑧", {"font": MED, "color": BLUE}), (" the limits: one deadline, N turns, a token reservation", {})],
                               [("⑤", {"font": MED, "color": BLUE}), (" the thought: one model turn", {})],
                               [("⑥ ⑦", {"font": MED, "color": BLUE}), (" classify the turn: tool calls go back to ③, anything else must be the answer", {})],
                               [("③", {"font": MED, "color": BLUE}), (" the action and observation", {})]], size=11, space=8)

# ------------------------------------------------------------------ 29 act step
s = slide(CONTENT, "Zoom in on ③: the act step",
          "Take the first pending call but leave it in the list until it is done, so a crash leaves it visible. Fingerprint "
          "it as tool name plus canonical arguments; a repeat fails with repeated_action before it runs. Execute with the "
          "application's context, never the model's. Append the observation tagged with the same call ID, then pop and "
          "save after every call. Order matters for crashes: a crash after execute re-runs the call on resume; harmless for "
          "reads, which is why Session 02 adds idempotent effects for writes.")
code(s, 0.4, 1.0, 5.6, 3.3, "while run.pending:\n    call = run.pending[0]\n    key = json.dumps([call[\"name\"], call[\"args\"]],\n"
     "                     sort_keys=True)\n    if key in seen:\n        raise ValueError(\"repeated_action\")\n    seen.add(key)\n"
     "    result = await execute(call[\"name\"], call[\"args\"],\n                           run.context)\n"
     "    run.messages.append({\"role\": \"tool\",\n        \"tool_call_id\": call[\"id\"], \"content\": json.dumps(result)})\n"
     "    run.pending.pop(0)\n    save()", size=9.5, hl=(10, 11))
text(s, 6.25, 1.0, 3.25, 3.4, [[("Leave it pending", {"font": MED, "color": NAVY})], "until done, so a crash shows what was unfinished.",
                               [("Refuse repeats", {"font": MED, "color": NAVY})], "same tool and arguments → repeated_action, before it runs.",
                               [("Run with app context", {"font": MED, "color": NAVY})], "identity and permissions come from Context, not the model.",
                               [("Tag the observation", {"font": MED, "color": NAVY})], "same call ID, or the conversation is invalid."], size=10.5, space=3)

# ------------------------------------------------------------------ 30 answer check
s = slide(CONTENT, "Zoom in on ⑦: the answer check",
          "Parse: the text must be bare JSON; prose or a code fence fails here without a named reason, the gap noted earlier. "
          "Shape: status, cause, summary and evidence must exist, or invalid_final_answer. Evidence: every cited source must "
          "appear in the tool results, or unobserved_citation; an empty list passes, which is right for insufficient "
          "evidence. Only then is the run completed. Ask: the model cites rb-rollback-v9, which no tool returned. The run "
          "fails with unobserved_citation. Later, security.validate_answer makes this stricter.")
code(s, 0.4, 1.0, 5.6, 3.0, "answer = json.loads(reply.message.get(\"content\") or \"{}\")\nif not {\"status\", \"cause\", \"summary\",\n"
     "        \"evidence\"} <= answer.keys():\n    raise ValueError(\"invalid_final_answer\")\n"
     "observed = \" \".join(m[\"content\"] for m in run.messages\n                    if m[\"role\"] == \"tool\")\n"
     "if any(src not in observed for src in answer[\"evidence\"]):\n    raise ValueError(\"unobserved_citation\")\n"
     "run.answer = answer\nrun.status = \"completed\"", size=9.5, hl=(4, 5, 6, 7))
text(s, 6.25, 1.0, 3.25, 3.4, [[("Parse", {"font": MED, "color": NAVY})], "bare JSON only; anything else fails unnamed.",
                               [("Shape", {"font": MED, "color": NAVY})], "four fields, or invalid_final_answer.",
                               [("Evidence", {"font": MED, "color": NAVY})], "every citation observed, or unobserved_citation.",
                               [("Accept", {"font": MED, "color": NAVY})], "only then: completed and saved."], size=10.5, space=3)
text(s, 0.4, 4.25, 9.0, 0.45, [[("Ask: ", {"font": MED, "color": NAVY}),
                                ("the model cites rb-rollback-v9, which no tool returned. What happens?", {})]], size=11.5)

# ------------------------------------------------------------------ 31 steps 2 and 4
s = slide(CONTENT, "Steps 2 and 4: providers, guards",
          "Step 2: paste the given MockProvider unchanged and build HTTPProvider for OpenAI-compatible and Anthropic, checked "
          "against fake HTTP responses, so no paid calls. Step 4: keep the step guard and add token, time and repeated-action "
          "guards with exact error codes. After step 3 the happy path, step limit and adapter tests pass; after step 4 all "
          "seven do.")
card(s, 0.4, 1.05, 4.45, 3.1, "Step 2 · providers.py",
     "Paste the given MockProvider unchanged. Build HTTPProvider(kind, model, api_key, base_url, transport) for OpenAI-compatible "
     "and Anthropic, normalized into Reply. Check both request shapes against fake HTTP responses: no paid calls.",
     "MODEL_PROVIDER defaults to mock")
card(s, 5.05, 1.05, 4.45, 3.1, "Step 4 · the guards",
     "Keep the step guard; add the token, time and repeated-action guards. Estimate input as serialized length ÷ 3, add the "
     "reply reservation, and refuse the call before it would exceed the budget. The deadline cancels a call in progress.",
     "step_budget · token_budget · TimeoutError · repeated_action")
text(s, 0.4, 4.35, 9.0, 0.4, [[("Checks: ", {"font": MED, "color": NAVY}), ("after step 3, 5 of 7 pass; after step 4, all 7.", {})]], size=11.5)

# ------------------------------------------------------------------ 32 break it
s = slide(CONTENT, "Break it on purpose",
          "Change one limit at a time, write down the failure you expect, then run it. In class we predict together and I "
          "run it; after class you do it in your own project with a throwaway script. Restore the normal limits and confirm "
          "the cited answer comes back.")
pr = [("Limits(steps=1)", "step_budget", "One step cannot both call a tool and answer."),
      ("Limits(tokens=1)", "token_budget", "Refused before the first model call."),
      ("seconds=0.001 + slow mock", "TimeoutError", "The whole run has one deadline."),
      ("MockProvider(repeat=True)", "repeated_action", "Same tool, same arguments, twice.")]
for i, (a, b, c) in enumerate(pr):
    x = 0.4 + i * 2.33
    box(s, x, 1.2, 2.13, 2.7)
    text(s, x + 0.1, 1.35, 1.98, 2.5, [[(a, {"font": MONO, "size": 8.5, "color": NAVY})], " ",
                                        [("→ " + b, {"font": MONO, "size": 12, "color": BLUE})], " ", c], size=11, space=4)
text(s, 0.4, 4.15, 9.0, 0.45, "Predict, run, compare. Then restore the normal limits and confirm the cited answer comes back.", size=11.5, color=GREY)

# ------------------------------------------------------------------ 33 evidence
s = slide(CONTENT, "Evidence, not claims",
          "Your project is done when the Session 01 check pack passes: seven tests, offline, in under a second, covering the "
          "cited happy path, each of the three limits, and both provider request shapes against fake HTTP. Then the verify "
          "prompt commits and pushes to GitLab. Unrun or skipped checks are not passes: look at the actual output, not the "
          "agent's summary. Run IDs and timings vary; status, cause and evidence must not.")
box(s, 0.4, 1.05, 3.6, 3.5)
text(s, 0.6, 1.2, 3.3, 1.4, [[("7 passed", {"font": "Zalando Sans ExtraLight", "size": 50, "color": BLUE})]])
text(s, 0.6, 2.45, 3.25, 2.0, [[("LESSON = 1 · pushed to GitLab", {"font": MED, "size": 12, "color": NAVY})],
                               "The cited happy path, the three limits and both provider request shapes, offline.",
                               [("Unrun or skipped checks are not passes.", {"color": GREY})]], size=11, space=6)
code(s, 4.3, 1.05, 5.2, 3.5, "{\n  \"capability\": \"bounded ReAct loop\",\n  \"answer\": {\n    \"status\": \"complete\",\n"
     "    \"cause\": \"connection_pool_exhaustion\",\n    \"summary\": \"INC-001: connection_pool_exhaustion\",\n"
     "    \"evidence\": [\"log-astro-001\", \"rb-pool-v1\"]\n  },\n  \"model_steps\": 3,\n  \"repeated_action\": \"repeated_action\"\n}",
     size=10, hl=(4, 6))

# ------------------------------------------------------------------ 34 takeaways
s = slide(CONTENT, "You own the loop",
          "Four takeaways. An agent is a ReAct loop. Authority is code. Every run is bounded and every stop has a name. "
          "Completed is not solved: look for cited evidence or an honest insufficient evidence. Keep your blueprint; later "
          "the console shows these same run events as a timeline.")
tk = [("An agent is a loop", "the model proposes, the runtime decides and executes, the tool observes."),
      ("Authority is code", "tenant, role and permitted tools are set by the application, never by model output."),
      ("Every run is bounded", "and every stop has a name and a reason."),
      ("Completed is not solved", "look for cited evidence or an honest \"insufficient evidence\".")]
for i, (h, b) in enumerate(tk):
    x, y = 0.4 + (i % 2) * 4.65, 1.05 + (i // 2) * 1.75
    box(s, x, y, 4.45, 1.55)
    text(s, x + 0.2, y + 0.12, 0.8, 0.6, [[(f"0{i+1}", {"font": "Zalando Sans ExtraLight", "size": 26, "color": BLUE})]])
    text(s, x + 1.0, y + 0.2, 3.3, 1.3, [[(h, {"font": MED, "size": 14, "color": NAVY})], b], size=11.5)

# ------------------------------------------------------------------ 35 after class
s = slide(CONTENT, "After class, and next session",
          "Run today's prompts in your own project, in order: the lesson-00 prompt, add the check pack, build steps 1 to 4, "
          "break it on purpose, verify and push, then review. Next session: reliable tools and MCP integration, starting "
          "from your Session 01 checkpoint. Today the model is a scripted mock; a few sessions from now we point the same "
          "platform at a real model by configuration alone.")
after = ["The lesson-00 prompt: your project, the course repository, your GitLab repository",
         "Add the check pack", "Build steps 1 to 4, one at a time, reading each change", "Break it on purpose",
         "Verify and push, then review"]
for i, a in enumerate(after):
    y = 1.05 + i * 0.62
    box(s, 0.4, y, 0.45, 0.45, fill=BLUE, line=None, shape=MSO_SHAPE.OVAL)
    text(s, 0.4, y, 0.45, 0.45, [[(str(i + 1), {"font": MED, "size": 12, "color": "FFFFFF"})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.0, y + 0.06, 4.6, 0.5, a, size=11.5)
box(s, 5.9, 1.05, 3.6, 3.0, fill=NAVY, line=None)
text(s, 6.1, 1.2, 3.25, 2.8, [[("Next session", {"font": MONO, "size": 9.5, "color": MINT})],
                              [("Reliable tools and MCP integration", {"font": MED, "size": 15, "color": "FFFFFF"})],
                              [("Validated tool contracts, application-owned permissions, an approval gate for a mock remediation, "
                                "and real MCP transport. Starts from your Session 01 checkpoint.", {"color": "C9D3F5"})]], size=11, space=8)
text(s, 0.4, 4.3, 9.0, 0.4, [[(PAGE, {"font": MONO, "size": 10, "color": BRIGHT})]])

# ------------------------------------------------------------------ 36 thank you
slide("Thank you slide 2", None, "Thank you. Questions.")

# drop the template's sample slides
for sid in TEMPLATE_SLIDES:
    prs.part.drop_rel(sid.get(qn("r:id")))
    prs.slides._sldIdLst.remove(sid)
used = {id(sl.slide_layout._element) for sl in prs.slides}
for m in prs.slide_masters:
    lst = m._element.find(qn("p:sldLayoutIdLst"))
    keep_one = None
    for lid in list(lst):
        layout = m.part.related_part(lid.get(qn("r:id"))).slide_layout
        if id(layout._element) in used:
            continue
        if keep_one is None and not any(id(l._element) in used for l in m.slide_layouts):
            keep_one = lid  # a master must keep at least one layout
            continue
        m.part.drop_rel(lid.get(qn("r:id")))
        lst.remove(lid)
prs.save(OUT)
print("slides:", len(prs.slides))
